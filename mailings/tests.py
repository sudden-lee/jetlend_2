import logging
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from inspect import GEN_CLOSED, getgeneratorstate
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Barrier
from traceback import format_exception
from unittest.mock import patch
from uuid import UUID, uuid4
from xml.etree.ElementTree import Element, ParseError, SubElement, fromstring, tostring
from zipfile import ZIP_DEFLATED, ZipFile

from defusedxml.common import DefusedXmlException
from django.contrib import admin
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import DatabaseError, IntegrityError, connections, transaction
from django.db.backends.utils import CursorWrapper
from django.db.models.query import QuerySet
from django.test import RequestFactory, TestCase, TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.cell.read_only import ReadOnlyCell
from openpyxl.formula.translate import TranslatorError
from openpyxl.utils.exceptions import InvalidFileException

from mailings.importing import (
    COLUMNS,
    MAX_ARCHIVE_ENTRIES,
    MAX_COLUMNS,
    ImportDatabaseError,
    ImportStats,
    RowError,
    WorkbookError,
    import_mailings,
    parse_row,
    save_batch,
    validate_first_worksheet,
)
from mailings.management.commands.send_mailings import Command as SendCommand
from mailings.models import Mailing
from mailings.sending import (
    LEASE,
    MAX_ATTEMPTS,
    PermanentTransportError,
    SendDatabaseError,
    SendStats,
    claim_next,
    delivery_idempotency_key,
    retry_delay,
    send_email,
    send_mailings,
)


class ImportTests(TestCase):
    def setUp(self) -> None:
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.path = Path(directory.name) / "mailings.xlsx"

    def workbook(self, rows: list[list[object]], headers: tuple[str, ...] = COLUMNS) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        workbook.save(self.path)
        workbook.close()
        return self.path

    def test_import_and_reimport_preserve_first_payload_and_sent_status(self) -> None:
        rows = [
            ["first", 1, "one@example.com", "First subject", "First message"],
            ["first", 2, "two@example.com", "Changed", "Changed"],
            ["second", 3, "three@example.com", "Second", "Second"],
            ["first", 4, "four@example.com", "Changed again", "Changed again"],
        ]
        self.assertEqual(
            import_mailings(self.workbook(rows), batch_size=2), ImportStats(4, 2, 2, 0)
        )
        first = Mailing.objects.get(external_id="first")
        self.assertEqual((first.user_id, first.subject), (1, "First subject"))
        first.status = Mailing.Status.SENT
        first.sent_at = timezone.now()
        first.save(update_fields=["status", "sent_at"])
        self.assertEqual(import_mailings(self.path, batch_size=1), ImportStats(4, 0, 4, 0))
        first.refresh_from_db()
        self.assertEqual(first.status, Mailing.Status.SENT)
        self.assertEqual(Mailing.objects.count(), 2)

    def test_duplicates_preserve_fifo_creation_order(self) -> None:
        rows = [
            ["first", 1, "user@example.com", "First", "Body"],
            ["second", 2, "user@example.com", "Second", "Body"],
            ["first", 3, "user@example.com", "Replacement", "Body"],
        ]
        self.assertEqual(import_mailings(self.workbook(rows)), ImportStats(3, 2, 1, 0))
        self.assertEqual(
            list(Mailing.objects.order_by("pk").values_list("external_id", "subject")),
            [("first", "First"), ("second", "Second")],
        )

    def test_reordered_headers_extra_columns_and_empty_rows(self) -> None:
        headers = ("message", "ignored", " email ", "user_id", "subject", "external_id")
        rows = [[None] * 6, ["Body", "extra", " user@example.com ", "123", "Subject", " id "]]
        self.assertEqual(import_mailings(self.workbook(rows, headers)), ImportStats(1, 1, 0, 0))
        mailing = Mailing.objects.get()
        self.assertEqual((mailing.external_id, mailing.user_id), ("id", 123))
        self.assertEqual((mailing.email, mailing.message), ("user@example.com", "Body"))

    def test_numeric_external_ids_deduplicate_against_text(self) -> None:
        rows = [[42, 1, "user@example.com", "Subject", "Body"]] * 2
        rows[1] = ["42", 2, "user@example.com", "Subject", "Body"]
        self.assertEqual(import_mailings(self.workbook(rows)), ImportStats(2, 1, 1, 0))
        self.assertEqual(Mailing.objects.get().external_id, "42")

    def test_large_user_id_stored_as_text_is_exact(self) -> None:
        row = ["id", "9223372036854775807", "user@example.com", "Subject", "Body"]
        self.assertEqual(import_mailings(self.workbook([row])), ImportStats(1, 1, 0, 0))
        self.assertEqual(Mailing.objects.get().user_id, 2**63 - 1)

    def test_invalid_rows_do_not_discard_valid_rows(self) -> None:
        valid = ["good", 1, "user@example.com", "Subject", "Body"]
        invalid = [
            ["bad-email", 1, "invalid", "Subject", "Body"],
            ["formula", 1, "user@example.com", "=1+1", "Body"],
            ["excel-error", 1, "user@example.com", "#DIV/0!", "Body"],
            ["boolean", True, "user@example.com", "Subject", "Body"],
            ["fractional", 1.5, "user@example.com", "Subject", "Body"],
            ["missing", 1, "user@example.com", "Subject", None],
            ["zero", 0, "user@example.com", "Subject", "Body"],
            ["overflow", "9223372036854775808", "user@example.com", "Subject", "Body"],
            ["header-injection", 1, "user@example.com", "Subject\nBcc: other", "Body"],
        ]
        output = StringIO()
        with self.assertLogs("mailings.importing", level="WARNING"):
            with self.assertRaises(CommandError):
                call_command(
                    "import_mailings", str(self.workbook([valid, *invalid])), stdout=output
                )
        self.assertEqual(Mailing.objects.count(), 1)
        self.assertIn(
            "Обработано: 10; создано: 1; пропущено: 0; ошибочных строк: 9", output.getvalue()
        )

    def test_parser_rejects_types_and_length_without_excel_coercion(self) -> None:
        invalid_values = {
            "external_id": [True, 1.5, 10**15, " " * 3, "a" * 256],
            "user_id": [True, 1.0, -1, "-1", "١", 10**15],
            "email": [10, "user@example.com\n", "a" * 255],
            "subject": [10, " ", "a" * 256],
            "message": [10, " ", "a" * 32768],
        }
        workbook = Workbook()
        sheet = workbook.active
        indexes = {name: index for index, name in enumerate(COLUMNS)}
        for name, values in invalid_values.items():
            for value in values:
                with self.subTest(column=name, value=repr(value)[:40]):
                    row = ["id", 1, "user@example.com", "Subject", "Body"]
                    row[indexes[name]] = value
                    cells = tuple(
                        ReadOnlyCell(sheet, 1, index, entry, data_type="s")
                        for index, entry in enumerate(row, start=1)
                    )
                    with self.assertRaises(RowError):
                        parse_row(cells, indexes)
        workbook.close()

    def test_empty_or_invalid_headers_reject_workbook(self) -> None:
        for headers in [(), COLUMNS[:-1], (*COLUMNS, "external_id")]:
            with self.subTest(headers=headers):
                with self.assertRaises(WorkbookError) as caught:
                    import_mailings(self.workbook([], headers))
                self.assertEqual(caught.exception.stats, ImportStats())
        self.assertFalse(Mailing.objects.exists())

    def test_missing_non_xlsx_and_corrupt_files_are_reported(self) -> None:
        for path in [self.path, self.path.with_suffix(".csv")]:
            with self.subTest(path=path), self.assertRaises(WorkbookError):
                import_mailings(path)
        self.path.write_bytes(b"not a ZIP archive")
        with self.assertRaises(WorkbookError) as caught:
            import_mailings(self.path)
        self.assertEqual(caught.exception.stats, ImportStats())

    def test_xml_entities_are_rejected_before_workbook_loading(self) -> None:
        original = self.workbook([]).read_bytes()
        malicious = b"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE Types [<!ENTITY payload "expanded">]>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
&payload;
</Types>"""
        with ZipFile(BytesIO(original)) as source, ZipFile(self.path, "w") as target:
            for entry in source.infolist():
                contents = (
                    malicious
                    if entry.filename == "[Content_Types].xml"
                    else source.read(entry.filename)
                )
                target.writestr(entry, contents)

        with self.assertRaises(WorkbookError) as caught:
            import_mailings(self.path)
        self.assertIsInstance(caught.exception.__cause__, DefusedXmlException)
        self.assertEqual(caught.exception.stats, ImportStats())
        self.assertFalse(Mailing.objects.exists())

    def test_oversized_workbook_metadata_is_rejected_before_loading(self) -> None:
        self.workbook([])
        with patch("mailings.importing.MAX_METADATA_BYTES", 1):
            with self.assertRaises(WorkbookError) as caught:
                import_mailings(self.path)
        self.assertIsInstance(caught.exception.__cause__, InvalidFileException)
        self.assertEqual(caught.exception.stats, ImportStats())

    def test_archive_with_too_many_parts_is_rejected_before_openpyxl(self) -> None:
        source = BytesIO()
        with ZipFile(source, "w", ZIP_DEFLATED) as archive:
            for number in range(MAX_ARCHIVE_ENTRIES + 1):
                archive.writestr(f"parts/{number}.xml", b"")
        source.seek(0)
        with self.assertRaisesRegex(InvalidFileException, "too many archive parts"):
            validate_first_worksheet(source)

    def test_missing_first_worksheet_is_rejected_instead_of_importing_second(self) -> None:
        workbook = Workbook()
        first = workbook.active
        first.append(COLUMNS)
        first.append(["first", 1, "one@example.com", "First", "Body"])
        second = workbook.create_sheet("second")
        second.append(COLUMNS)
        second.append(["second", 2, "two@example.com", "Second", "Body"])
        workbook.save(self.path)
        workbook.close()
        original = self.path.read_bytes()
        with ZipFile(BytesIO(original)) as source, ZipFile(self.path, "w") as target:
            for entry in source.infolist():
                if entry.filename != "xl/worksheets/sheet1.xml":
                    target.writestr(entry, source.read(entry.filename))
        with self.assertRaises(WorkbookError) as caught:
            import_mailings(self.path)
        self.assertEqual(caught.exception.stats, ImportStats())
        self.assertFalse(Mailing.objects.exists())

    def test_renamed_worksheet_part_follows_relative_and_absolute_relationships(self) -> None:
        original = self.workbook([["first", 1, "one@example.com", "First", "Body"]]).read_bytes()
        renamed = "xl/data/letters.xml"
        for relationship_target in ["data/letters.xml", "/xl/data/letters.xml"]:
            with self.subTest(target=relationship_target):
                Mailing.objects.all().delete()
                with ZipFile(BytesIO(original)) as source, ZipFile(self.path, "w") as target:
                    for entry in source.infolist():
                        contents = source.read(entry.filename)
                        name = entry.filename
                        if name == "xl/worksheets/sheet1.xml":
                            name = renamed
                        elif name == "xl/_rels/workbook.xml.rels":
                            relationships = fromstring(contents)
                            for relation in relationships:
                                if relation.get("Type", "").endswith("/worksheet"):
                                    relation.set("Target", relationship_target)
                            contents = tostring(relationships)
                        elif name == "[Content_Types].xml":
                            types = fromstring(contents)
                            for content_type in types:
                                if content_type.get("PartName") == "/xl/worksheets/sheet1.xml":
                                    content_type.set("PartName", f"/{renamed}")
                            contents = tostring(types)
                        target.writestr(name, contents)
                self.assertEqual(import_mailings(self.path), ImportStats(1, 1, 0, 0))
                self.assertEqual(Mailing.objects.get().external_id, "first")

    def test_shared_formula_read_failure_retains_valid_buffer(self) -> None:
        rows = [
            ["good", 1, "one@example.com", "First", "Body", None, None, "=A1"],
            ["broken", 2, "two@example.com", "Second", "Body", None, "=A1"],
        ]
        original = self.workbook(rows, (*COLUMNS, "extra1", "extra2", "extra3")).read_bytes()
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with ZipFile(BytesIO(original)) as source, ZipFile(self.path, "w") as target:
            for entry in source.infolist():
                contents = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    sheet = fromstring(contents)
                    first = sheet.find(f".//{{{namespace}}}c[@r='H2']/{{{namespace}}}f")
                    first.attrib.update(t="shared", si="0", ref="G2:H3")
                    second = sheet.find(f".//{{{namespace}}}c[@r='G3']/{{{namespace}}}f")
                    second.attrib.update(t="shared", si="0")
                    second.text = None
                    contents = tostring(sheet)
                target.writestr(entry, contents)
        with self.assertRaises(WorkbookError) as caught:
            import_mailings(self.path)
        self.assertIsInstance(caught.exception.__cause__, TranslatorError)
        self.assertEqual(caught.exception.stats, ImportStats(1, 1, 0, 0))
        self.assertEqual(Mailing.objects.get().external_id, "good")

    def test_overflowing_iso_duration_read_failure_retains_valid_buffer(self) -> None:
        rows = [
            ["good", 1, "one@example.com", "First", "Body"],
            ["broken", 2, "two@example.com", "Second", "Body"],
        ]
        original = self.workbook(rows).read_bytes()
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with ZipFile(BytesIO(original)) as source, ZipFile(self.path, "w") as target:
            for entry in source.infolist():
                contents = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    sheet = fromstring(contents)
                    cell = sheet.find(f".//{{{namespace}}}c[@r='B3']")
                    cell.clear()
                    cell.attrib.update(r="B3", t="d")
                    SubElement(cell, f"{{{namespace}}}v").text = "PT999999999999999999999999H"
                    contents = tostring(sheet)
                target.writestr(entry, contents)
        with self.assertRaises(WorkbookError) as caught:
            import_mailings(self.path)
        self.assertIsInstance(caught.exception.__cause__, OverflowError)
        self.assertEqual(caught.exception.stats, ImportStats(1, 1, 0, 0))
        self.assertEqual(Mailing.objects.get().external_id, "good")

    def test_import_command_reports_partial_database_failure_without_retry(self) -> None:
        rows = [
            ["first", 1, "one@example.com", "First", "Body"],
            ["second", 2, "two@example.com", "Second", "Body"],
        ]
        self.workbook(rows)
        real_execute = CursorWrapper.execute
        error = DatabaseError("private database payload")
        calls = 0

        def failing_execute(cursor: CursorWrapper, query: object, params: object = None) -> object:
            nonlocal calls
            calls += 1
            if calls == 2:
                raise error
            return real_execute(cursor, query, params)

        output = StringIO()
        with patch.object(CursorWrapper, "execute", new=failing_execute):
            with self.assertRaises(CommandError) as caught:
                call_command("import_mailings", str(self.path), batch_size=1, stdout=output)
        self.assertEqual(calls, 2)
        self.assertEqual(list(Mailing.objects.values_list("external_id", flat=True)), ["first"])
        self.assertIn(
            "Обработано: 2; создано: 1; пропущено: 0; ошибочных строк: 0; незавершённых строк: 1",
            output.getvalue(),
        )
        self.assertIsNone(caught.exception.__cause__)
        self.assertTrue(caught.exception.__suppress_context__)
        self.assertIsInstance(caught.exception.__context__, ImportDatabaseError)
        self.assertIs(caught.exception.__context__.__cause__, error)
        self.assertNotIn("private database payload", str(caught.exception))
        self.assertNotIn("private database payload", output.getvalue())
        self.assertNotIn("private database payload", "".join(format_exception(caught.exception)))

    def test_invalid_shared_string_reference_reports_partial_results(self) -> None:
        rows = [
            ["good", 1, "user@example.com", "Subject", "Body"],
            ["broken", 2, "user@example.com", "Subject", "Body"],
        ]
        original = self.workbook(rows).read_bytes()
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        content_namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
        with ZipFile(BytesIO(original)) as source, ZipFile(self.path, "w") as target:
            for entry in source.infolist():
                contents = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    sheet = fromstring(contents)
                    cell = sheet.find(f".//{{{namespace}}}c[@r='A3']")
                    self.assertIsNotNone(cell)
                    cell.clear()
                    cell.attrib.update(r="A3", t="s")
                    SubElement(cell, f"{{{namespace}}}v").text = "999"
                    contents = tostring(sheet)
                elif entry.filename == "[Content_Types].xml":
                    types = fromstring(contents)
                    SubElement(
                        types,
                        f"{{{content_namespace}}}Override",
                        PartName="/xl/sharedStrings.xml",
                        ContentType=(
                            "application/vnd.openxmlformats-officedocument."
                            "spreadsheetml.sharedStrings+xml"
                        ),
                    )
                    contents = tostring(types)
                target.writestr(entry, contents)
            strings = Element(f"{{{namespace}}}sst", count="0", uniqueCount="0")
            target.writestr("xl/sharedStrings.xml", tostring(strings))

        with self.assertRaises(WorkbookError) as caught:
            import_mailings(self.path)
        self.assertIsInstance(caught.exception.__cause__, IndexError)
        self.assertEqual(caught.exception.stats, ImportStats(1, 1, 0, 0))
        self.assertEqual(Mailing.objects.get().external_id, "good")
        output = StringIO()
        with self.assertRaises(CommandError):
            call_command("import_mailings", str(self.path), stdout=output)
        self.assertIn(
            "Обработано: 1; создано: 0; пропущено: 1; ошибочных строк: 0", output.getvalue()
        )

    def test_source_handle_is_closed_when_workbook_loading_fails(self) -> None:
        self.workbook([])
        source = self.path.open("rb")
        self.addCleanup(source.close)
        with patch.object(Path, "open", return_value=source):
            with patch(
                "mailings.importing.load_workbook",
                side_effect=ValueError("private workbook payload"),
            ):
                with self.assertRaises(WorkbookError):
                    import_mailings(self.path)
        self.assertTrue(source.closed)

        output = StringIO()
        with patch(
            "mailings.importing.load_workbook",
            side_effect=ValueError("private workbook payload"),
        ):
            with self.assertRaises(CommandError) as caught:
                call_command("import_mailings", str(self.path), stdout=output)
        self.assertNotIn("private workbook payload", "".join(format_exception(caught.exception)))

    def test_invalid_workbook_sheet_id_is_reported(self) -> None:
        original = self.workbook([]).read_bytes()
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with ZipFile(BytesIO(original)) as source, ZipFile(self.path, "w") as target:
            for entry in source.infolist():
                contents = source.read(entry.filename)
                if entry.filename == "xl/workbook.xml":
                    workbook = fromstring(contents)
                    sheet = workbook.find(f".//{{{namespace}}}sheet")
                    self.assertIsNotNone(sheet)
                    sheet.set("sheetId", "no")
                    contents = tostring(workbook)
                target.writestr(entry, contents)
        with self.assertRaises(WorkbookError) as caught:
            import_mailings(self.path)
        self.assertIsInstance(caught.exception.__cause__, TypeError)
        self.assertEqual(caught.exception.stats, ImportStats())
        output = StringIO()
        with self.assertRaises(CommandError):
            call_command("import_mailings", str(self.path), stdout=output)
        self.assertIn(
            "Обработано: 0; создано: 0; пропущено: 0; ошибочных строк: 0", output.getvalue()
        )

    def test_bad_headers_close_source_workbook_and_row_generator(self) -> None:
        self.workbook([["id", 1, "user@example.com", "Subject"]], headers=COLUMNS[:-1])
        source = self.path.open("rb")
        self.addCleanup(source.close)
        workbook = load_workbook(source, read_only=True)
        self.addCleanup(workbook.close)
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows()
        with patch.object(Path, "open", return_value=source):
            with patch("mailings.importing.load_workbook", return_value=workbook):
                with patch.object(sheet, "iter_rows", return_value=rows):
                    with self.assertRaises(WorkbookError):
                        import_mailings(self.path)
        self.assertTrue(source.closed)
        self.assertIsNone(workbook._archive.fp)
        self.assertEqual(getgeneratorstate(rows), GEN_CLOSED)

    def test_invalid_batch_size_is_rejected_before_reading(self) -> None:
        for size in [0, -1, 501]:
            with self.subTest(size=size), self.assertRaises(ValueError):
                import_mailings(self.path, batch_size=size)

    def test_unexpected_parser_errors_propagate_and_close_resources(self) -> None:
        for parser in ["parse_headers", "parse_row"]:
            with self.subTest(parser=parser):
                self.workbook([["id", 1, "user@example.com", "Subject", "Body"]])
                source = self.path.open("rb")
                self.addCleanup(source.close)
                workbook = load_workbook(source, read_only=True)
                self.addCleanup(workbook.close)
                sheet = workbook.worksheets[0]
                rows = sheet.iter_rows()
                error = TypeError("programming error")
                with (
                    patch.object(Path, "open", return_value=source),
                    patch("mailings.importing.load_workbook", return_value=workbook),
                    patch.object(sheet, "iter_rows", return_value=rows),
                    patch(f"mailings.importing.{parser}", side_effect=error),
                    patch("mailings.importing.save_batch") as save,
                ):
                    with self.assertRaises(TypeError) as caught:
                        import_mailings(self.path)
                self.assertIs(caught.exception, error)
                save.assert_not_called()
                self.assertTrue(source.closed)
                self.assertIsNone(workbook._archive.fp)
                self.assertEqual(getgeneratorstate(rows), GEN_CLOSED)

    def test_batch_errors_propagate_without_retry(self) -> None:
        self.workbook([["id", 1, "user@example.com", "Subject", "Body"]])
        error = ValueError("programming error")
        with patch("mailings.importing.save_batch", side_effect=error) as save:
            with self.assertRaises(ValueError) as caught:
                import_mailings(self.path, batch_size=1)
        self.assertIs(caught.exception, error)
        save.assert_called_once()

        error = DatabaseError("database error")
        with patch("mailings.importing.save_batch", side_effect=error) as save:
            with self.assertRaises(ImportDatabaseError) as caught:
                import_mailings(self.path, batch_size=1)
        self.assertIs(caught.exception.__cause__, error)
        self.assertEqual(caught.exception.stats, ImportStats(1, 0, 0, 0))
        save.assert_called_once()

    def test_read_failure_retains_read_rows_and_reports_partial_stats(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(COLUMNS)
        sheet.append(["id", 1, "user@example.com", "Subject", "Body"])
        headers, data = tuple(sheet[1]), tuple(sheet[2])
        self.path.write_bytes(b"mock workbook")

        def damaged_rows(*, max_col: int):
            self.assertEqual(max_col, MAX_COLUMNS)
            yield headers
            yield data
            raise ParseError("damaged worksheet")

        with (
            patch("mailings.importing.validate_first_worksheet"),
            patch("mailings.importing.load_workbook", return_value=workbook),
            patch.object(sheet, "reset_dimensions", create=True),
            patch.object(sheet, "iter_rows", side_effect=damaged_rows),
            self.assertRaises(WorkbookError) as caught,
        ):
            import_mailings(self.path)
        self.assertEqual(caught.exception.stats, ImportStats(1, 1, 0, 0))
        self.assertEqual(Mailing.objects.get().external_id, "id")

    def test_database_prevents_duplicate_external_id(self) -> None:
        data = dict(external_id="id", user_id=1, email="user@example.com", subject="S", message="B")
        Mailing.objects.create(**data)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Mailing.objects.create(**data)


class SendingTests(TestCase):
    def create_mailing(self, **changes: object) -> Mailing:
        data = {
            "external_id": str(uuid4()),
            "user_id": 1,
            "email": "private@example.com",
            "subject": "Private subject",
            "message": "Private body",
        }
        data.update(changes)
        return Mailing.objects.create(**data)

    def test_simulated_transport_has_required_delay_and_private_log(self) -> None:
        mailing = self.create_mailing()
        with patch("mailings.sending.random.randint", return_value=7) as randint:
            with patch("mailings.sending.time.sleep") as sleep:
                with self.assertLogs("mailings.sending", level="INFO") as logs:
                    send_email(mailing)
        randint.assert_called_once_with(5, 20)
        sleep.assert_called_once_with(7)
        key = delivery_idempotency_key(mailing)
        self.assertEqual(
            logs.output,
            [f"INFO:mailings.sending:Send EMAIL mailing_id={mailing.pk} idempotency_key={key}"],
        )
        self.assertEqual(len(key), 64)
        self.assertNotIn(mailing.external_id, key)

    def test_broken_transport_log_is_a_failed_send(self) -> None:
        class BrokenStream:
            def write(self, value: str) -> None:
                raise BrokenPipeError("private transport payload")

        mailing = self.create_mailing()
        output = StringIO()
        diagnostics = StringIO()
        handler = logging.getLogger("mailings").handlers[0]
        with (
            patch.object(handler, "stream", BrokenStream()),
            patch("sys.stderr", diagnostics),
            patch.object(logging, "raiseExceptions", True),
            patch("mailings.sending.time.sleep") as sleep,
            self.assertRaises(CommandError) as caught,
        ):
            call_command("send_mailings", stdout=output)
        sleep.assert_called_once()
        mailing.refresh_from_db()
        self.assertEqual(
            (mailing.status, mailing.last_error), (Mailing.Status.RETRYING, "BrokenPipeError")
        )
        self.assertIsNone(mailing.sent_at)
        self.assertIsNotNone(mailing.next_attempt_at)
        self.assertIn("Отправлено: 0; назначено повторов: 1", output.getvalue())
        self.assertNotIn(
            "private", output.getvalue() + diagnostics.getvalue() + str(caught.exception)
        )

    def test_closed_log_stream_does_not_hide_failed_send_report(self) -> None:
        mailing = self.create_mailing()
        output = StringIO()
        closed = StringIO()
        closed.close()
        handler = logging.getLogger("mailings").handlers[0]
        command = SendCommand()
        with (
            patch.object(handler, "stream", closed),
            patch("sys.stderr", closed),
            patch.object(logging, "raiseExceptions", True),
            patch("mailings.sending.time.sleep"),
            self.assertRaises(CommandError) as caught,
        ):
            call_command(command, stdout=output)
        mailing.refresh_from_db()
        self.assertEqual(
            (mailing.status, mailing.last_error), (Mailing.Status.RETRYING, "ValueError")
        )
        self.assertIsNotNone(mailing.next_attempt_at)
        self.assertIn("Отправлено: 0; назначено повторов: 1", output.getvalue())
        self.assertEqual(str(caught.exception), "Очередь обработана с ошибками.")

    def test_success_and_repeat_do_not_resend(self) -> None:
        mailing = self.create_mailing()
        with patch("mailings.sending.send_email") as transport:
            self.assertEqual(send_mailings(), SendStats(1, 0, 0))
            self.assertEqual(send_mailings(), SendStats())
        transport.assert_called_once()
        mailing.refresh_from_db()
        self.assertEqual((mailing.status, mailing.attempts), (Mailing.Status.SENT, 1))
        self.assertIsNotNone(mailing.sent_at)
        self.assertIsNone(mailing.claim_token)
        self.assertIsNone(mailing.claimed_at)

    def test_transient_failure_is_isolated_and_retried_when_due(self) -> None:
        first, second = self.create_mailing(), self.create_mailing()
        with patch("mailings.sending.send_email", side_effect=[RuntimeError("secret"), None]):
            with self.assertLogs("mailings.sending", level="WARNING") as logs:
                self.assertEqual(send_mailings(), SendStats(sent=1, retried=1))
        self.assertNotIn("secret", " ".join(logs.output))
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual(
            (first.status, first.last_error), (Mailing.Status.RETRYING, "RuntimeError")
        )
        self.assertGreater(first.next_attempt_at, timezone.now())
        self.assertEqual(second.status, Mailing.Status.SENT)
        with patch("mailings.sending.send_email") as transport:
            self.assertEqual(send_mailings(), SendStats())
            transport.assert_not_called()
            Mailing.objects.filter(pk=first.pk).update(next_attempt_at=timezone.now())
            self.assertEqual(send_mailings(), SendStats(sent=1))
        first.refresh_from_db()
        self.assertEqual(
            (first.status, first.attempts, first.last_error), (Mailing.Status.SENT, 2, "")
        )
        self.assertIsNone(first.next_attempt_at)

    def test_expired_claim_is_recovered_while_active_claim_is_skipped(self) -> None:
        now = timezone.now()
        active = self.create_mailing(
            status=Mailing.Status.PROCESSING, claimed_at=now, claim_token=uuid4()
        )
        expired = self.create_mailing(
            status=Mailing.Status.PROCESSING,
            claimed_at=now - LEASE - timedelta(seconds=1),
            claim_token=uuid4(),
            attempts=1,
        )
        with patch("mailings.sending.send_email") as transport:
            self.assertEqual(send_mailings(), SendStats(1, 0, 0))
        self.assertEqual(transport.call_args.args[0].pk, expired.pk)
        active.refresh_from_db()
        expired.refresh_from_db()
        self.assertEqual(active.status, Mailing.Status.PROCESSING)
        self.assertEqual((expired.status, expired.attempts), (Mailing.Status.SENT, 2))

    def test_retry_limit_does_not_change_unclaimed_failed_mailings(self) -> None:
        first = self.create_mailing(
            status=Mailing.Status.FAILED, attempts=1, last_error="RuntimeError"
        )
        second = self.create_mailing(
            status=Mailing.Status.FAILED, attempts=2, last_error="TimeoutError"
        )
        with patch("mailings.sending.send_email") as transport:
            self.assertEqual(send_mailings(limit=1, retry_failed=True), SendStats(1, 0, 0))
        transport.assert_called_once()
        self.assertEqual(transport.call_args.args[0].pk, first.pk)
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual((first.status, first.attempts), (Mailing.Status.SENT, 2))
        self.assertEqual(
            (second.status, second.attempts, second.last_error, second.claim_token),
            (Mailing.Status.FAILED, 2, "TimeoutError", None),
        )
        self.assertIsNone(second.claimed_at)

    def test_retry_failure_is_attempted_once_per_mailing_per_run(self) -> None:
        mailings = [
            self.create_mailing(status=Mailing.Status.FAILED),
            self.create_mailing(status=Mailing.Status.FAILED),
            self.create_mailing(),
        ]
        # A guard limit makes a broken retry loop fail promptly rather than hang the test.
        with patch("mailings.sending.send_email", side_effect=RuntimeError("secret")) as transport:
            with self.assertLogs("mailings.sending", level="WARNING"):
                self.assertEqual(send_mailings(limit=4, retry_failed=True), SendStats(retried=3))
        self.assertEqual(
            [call.args[0].pk for call in transport.call_args_list],
            [mailing.pk for mailing in mailings],
        )
        for mailing in mailings:
            mailing.refresh_from_db()
            self.assertEqual((mailing.status, mailing.attempts), (Mailing.Status.RETRYING, 1))
            self.assertIsNotNone(mailing.next_attempt_at)

    def test_claim_returns_updated_fields(self) -> None:
        mailing = self.create_mailing(attempts=4, last_error="PreviousError")
        claimed = claim_next(mailing.pk)
        self.assertIsNotNone(claimed)
        mailing.refresh_from_db()
        fields = ["status", "attempts", "claimed_at", "claim_token", "last_error"]
        self.assertEqual(
            [getattr(claimed, field) for field in fields],
            [getattr(mailing, field) for field in fields],
        )
        self.assertEqual((mailing.status, mailing.attempts), (Mailing.Status.PROCESSING, 5))
        self.assertIsNotNone(mailing.claimed_at)
        self.assertIsNotNone(mailing.claim_token)
        self.assertEqual(mailing.last_error, "")

    def test_lost_claim_cannot_finish_success_or_failure(self) -> None:
        for fail in [False, True]:
            with self.subTest(fail=fail):
                mailing = self.create_mailing()
                new_token = uuid4()

                def steal_claim(
                    claimed: Mailing,
                    *,
                    idempotency_key: str,
                    token: UUID = new_token,
                    should_fail: bool = fail,
                ) -> None:
                    self.assertEqual(idempotency_key, delivery_idempotency_key(claimed))
                    Mailing.objects.filter(pk=claimed.pk).update(claim_token=token)
                    if should_fail:
                        raise RuntimeError("transport failed after lease was lost")

                with patch("mailings.sending.send_email", side_effect=steal_claim):
                    with self.assertLogs("mailings.sending", level="WARNING"):
                        self.assertEqual(send_mailings(limit=1), SendStats(0, 0, 1))
                mailing.refresh_from_db()
                self.assertEqual(mailing.status, Mailing.Status.PROCESSING)
                self.assertEqual(mailing.claim_token, new_token)
                self.assertIsNone(mailing.sent_at)
                self.assertEqual(mailing.last_error, "")

    def test_future_retry_is_not_claimed_until_due(self) -> None:
        mailing = self.create_mailing(
            status=Mailing.Status.RETRYING,
            attempts=1,
            next_attempt_at=timezone.now() + timedelta(minutes=1),
        )
        self.assertIsNone(claim_next(mailing.pk))
        Mailing.objects.filter(pk=mailing.pk).update(next_attempt_at=timezone.now())
        claimed = claim_next(mailing.pk)
        self.assertIsNotNone(claimed)
        self.assertEqual((claimed.pk, claimed.attempts), (mailing.pk, 2))

    def test_permanent_failure_is_terminal_without_retry(self) -> None:
        mailing = self.create_mailing()
        with patch("mailings.sending.send_email", side_effect=PermanentTransportError("private")):
            with self.assertLogs("mailings.sending", level="ERROR") as logs:
                self.assertEqual(send_mailings(), SendStats(failed=1))
        self.assertNotIn("private", " ".join(logs.output))
        mailing.refresh_from_db()
        self.assertEqual(
            (mailing.status, mailing.attempts, mailing.last_error),
            (Mailing.Status.FAILED, 1, "PermanentTransportError"),
        )
        self.assertIsNone(mailing.next_attempt_at)

    def test_expired_last_attempt_becomes_terminal_without_transport(self) -> None:
        mailing = self.create_mailing(
            status=Mailing.Status.PROCESSING,
            attempts=MAX_ATTEMPTS,
            claimed_at=timezone.now() - LEASE - timedelta(seconds=1),
            claim_token=uuid4(),
        )
        with patch("mailings.sending.send_email") as transport:
            self.assertEqual(send_mailings(), SendStats(failed=1))
        transport.assert_not_called()
        mailing.refresh_from_db()
        self.assertEqual(
            (mailing.status, mailing.attempts, mailing.last_error),
            (Mailing.Status.FAILED, MAX_ATTEMPTS, "AttemptsExhausted"),
        )

    def test_retry_delay_is_exponential_jittered_and_capped(self) -> None:
        with patch("mailings.sending.random.randint", side_effect=lambda low, high: high):
            self.assertEqual(retry_delay(1), timedelta(seconds=37))
            self.assertEqual(retry_delay(99), timedelta(seconds=3600))

    def test_limit_and_id_cutoff_bound_delivery(self) -> None:
        first, second = self.create_mailing(), self.create_mailing()
        with patch(
            "mailings.sending.send_email",
            side_effect=lambda _mailing, **_kwargs: self.create_mailing(),
        ):
            self.assertEqual(send_mailings(limit=1), SendStats(1, 0, 0))
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual(first.status, Mailing.Status.SENT)
        self.assertEqual(second.status, Mailing.Status.PENDING)
        with patch(
            "mailings.sending.send_email",
            side_effect=lambda _mailing, **_kwargs: self.create_mailing(),
        ):
            self.assertEqual(send_mailings(), SendStats(2, 0, 0))
        self.assertEqual(Mailing.objects.filter(status=Mailing.Status.PENDING).count(), 2)

    def test_empty_queue_and_invalid_limit(self) -> None:
        self.assertEqual(send_mailings(), SendStats())
        for limit in [0, -1]:
            with self.subTest(limit=limit), self.assertRaises(ValueError):
                send_mailings(limit=limit)

    def test_database_failure_after_transport_reports_unconfirmed_result(self) -> None:
        mailing = self.create_mailing()
        real_update = QuerySet.update
        error = DatabaseError("private database payload")

        def fail_final_update(queryset: QuerySet, **values: object) -> int:
            if values.get("status") == Mailing.Status.SENT:
                raise error
            return real_update(queryset, **values)

        output = StringIO()
        with (
            patch.object(QuerySet, "update", new=fail_final_update),
            patch("mailings.sending.send_email") as transport,
            self.assertRaises(CommandError) as caught,
        ):
            call_command("send_mailings", stdout=output)
        transport.assert_called_once()
        mailing.refresh_from_db()
        self.assertEqual((mailing.status, mailing.attempts), (Mailing.Status.PROCESSING, 1))
        self.assertIsNotNone(mailing.claim_token)
        self.assertIsNone(mailing.sent_at)
        self.assertIsNone(caught.exception.__cause__)
        self.assertTrue(caught.exception.__suppress_context__)
        typed = caught.exception.__context__
        self.assertIsInstance(typed, SendDatabaseError)
        self.assertIs(typed.__cause__, error)
        self.assertEqual((typed.stats, typed.unconfirmed), (SendStats(), 1))
        self.assertIn("неподтверждённых результатов: 1", output.getvalue())
        self.assertNotIn("private", output.getvalue() + str(caught.exception))
        self.assertNotIn("private database payload", "".join(format_exception(caught.exception)))

    def test_database_failure_before_transport_reports_zero_unconfirmed(self) -> None:
        self.create_mailing()
        for point in ["aggregate", "claim"]:
            with self.subTest(point=point):
                error = DatabaseError("private database payload")
                failure = (
                    patch.object(QuerySet, "aggregate", side_effect=error)
                    if point == "aggregate"
                    else patch("mailings.sending.claim_next", side_effect=error)
                )
                output = StringIO()
                with (
                    failure,
                    patch("mailings.sending.send_email") as transport,
                    self.assertRaises(CommandError) as caught,
                ):
                    call_command("send_mailings", stdout=output)
                transport.assert_not_called()
                self.assertIsNone(caught.exception.__cause__)
                self.assertTrue(caught.exception.__suppress_context__)
                typed = caught.exception.__context__
                self.assertIsInstance(typed, SendDatabaseError)
                self.assertIs(typed.__cause__, error)
                self.assertEqual((typed.stats, typed.unconfirmed), (SendStats(), 0))
                self.assertIn("неподтверждённых результатов: 0", output.getvalue())
                self.assertNotIn("private", output.getvalue() + str(caught.exception))
                self.assertNotIn(
                    "private database payload", "".join(format_exception(caught.exception))
                )

    def test_send_command_reports_failure_and_success(self) -> None:
        self.create_mailing()
        with patch("mailings.sending.send_email", side_effect=RuntimeError("secret")):
            with self.assertLogs("mailings.sending", level="WARNING"):
                with self.assertRaises(CommandError):
                    call_command("send_mailings", stdout=StringIO())
        output = StringIO()
        with patch("mailings.sending.send_email"):
            Mailing.objects.update(next_attempt_at=timezone.now())
            call_command("send_mailings", stdout=output)
        self.assertIn(
            "Отправлено: 1; назначено повторов: 0; окончательных ошибок: 0",
            output.getvalue(),
        )


class PostgresConcurrencyTests(TransactionTestCase):
    def test_concurrent_duplicate_batches_keep_exact_counters(self) -> None:
        barrier = Barrier(2)

        def insert(subject: str) -> ImportStats:
            connection = connections["default"]
            try:
                stats = ImportStats(processed=1)
                mailing = Mailing(
                    external_id="same-id",
                    user_id=1,
                    email="user@example.com",
                    subject=subject,
                    message="Body",
                )
                barrier.wait(timeout=5)
                save_batch([mailing], stats)
                return stats
            finally:
                connection.close()

        with ThreadPoolExecutor(max_workers=2) as executor:
            futures = [executor.submit(insert, subject) for subject in ("First", "Second")]
            results = [future.result(timeout=10) for future in futures]

        counters = sorted((result.created, result.skipped) for result in results)
        self.assertEqual(counters, [(0, 1), (1, 0)])
        self.assertEqual(Mailing.objects.filter(external_id="same-id").count(), 1)

    def test_concurrent_workers_claim_different_mailings(self) -> None:
        mailings = [
            Mailing.objects.create(
                external_id=f"worker-{number}",
                user_id=number,
                email=f"worker-{number}@example.com",
                subject="Subject",
                message="Body",
            )
            for number in (1, 2)
        ]
        barrier = Barrier(2)
        max_id = mailings[-1].pk

        def claim() -> int | None:
            connection = connections["default"]
            try:
                barrier.wait(timeout=5)
                mailing = claim_next(max_id)
                return mailing.pk if mailing else None
            finally:
                connection.close()

        with ThreadPoolExecutor(max_workers=2) as executor:
            claimed_ids = [
                future.result(timeout=10) for future in [executor.submit(claim) for _ in range(2)]
            ]

        self.assertCountEqual(claimed_ids, [mailing.pk for mailing in mailings])
        self.assertEqual(
            Mailing.objects.filter(
                status=Mailing.Status.PROCESSING,
                claim_token__isnull=False,
                claimed_at__isnull=False,
            ).count(),
            2,
        )


class AdminTests(TestCase):
    @classmethod
    def setUpTestData(cls) -> None:
        cls.staff = get_user_model().objects.create_superuser(username="admin")
        cls.mailing = Mailing.objects.create(
            external_id="admin-visible",
            user_id=42,
            email="admin@example.com",
            subject="Visible subject",
            message="Visible message",
        )

    def test_model_has_localized_names_and_string_representation(self) -> None:
        self.assertEqual(str(self.mailing), "admin-visible")
        self.assertEqual(Mailing._meta.verbose_name, "рассылка")
        self.assertEqual(Mailing._meta.verbose_name_plural, "рассылки")
        self.assertEqual(Mailing._meta.get_field("email").verbose_name, "email получателя")
        self.assertEqual(self.mailing.get_status_display(), "Ожидает отправки")

    def test_admin_changelist_is_registered_localized_and_searchable(self) -> None:
        changelist_url = reverse("admin:mailings_mailing_changelist")
        anonymous_response = self.client.get(changelist_url)
        self.assertEqual(anonymous_response.status_code, 302)
        self.assertTrue(anonymous_response.url.startswith(f"{reverse('admin:login')}?next="))

        self.client.force_login(self.staff)
        for query in (self.mailing.external_id, "42", "admin@example", "Visible subject"):
            with self.subTest(query=query):
                response = self.client.get(changelist_url, {"q": query})
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, self.mailing.external_id)

        response = self.client.get(changelist_url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Управление рассылками")
        self.assertContains(response, self.mailing.external_id)
        self.assertContains(response, self.mailing.email)
        self.assertContains(response, "Ожидает отправки")

    def test_admin_forbids_bypassing_queue_workflow(self) -> None:
        model_admin = admin.site._registry[Mailing]
        request = RequestFactory().get("/admin/mailings/mailing/")
        request.user = self.staff
        self.assertFalse(model_admin.has_add_permission(request))
        self.assertFalse(model_admin.has_change_permission(request, self.mailing))
        self.assertFalse(model_admin.has_delete_permission(request, self.mailing))

    def test_database_rejects_zero_user_id(self) -> None:
        with self.assertRaises(IntegrityError), transaction.atomic():
            Mailing.objects.create(
                external_id="zero-user",
                user_id=0,
                email="zero@example.com",
                subject="Subject",
                message="Body",
            )

    def test_database_rejects_inconsistent_status_fields(self) -> None:
        with self.assertRaises(IntegrityError), transaction.atomic():
            Mailing.objects.create(
                external_id="invalid-sent",
                user_id=1,
                email="invalid@example.com",
                subject="Subject",
                message="Body",
                status=Mailing.Status.SENT,
            )
        with self.assertRaises(IntegrityError), transaction.atomic():
            Mailing.objects.create(
                external_id="invalid-retry",
                user_id=1,
                email="invalid@example.com",
                subject="Subject",
                message="Body",
                status=Mailing.Status.RETRYING,
            )
