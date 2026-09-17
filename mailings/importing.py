import logging
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from mailings.models import Mailing

logger = logging.getLogger(__name__)

COLUMNS = ("external_id", "user_id", "email", "subject", "message")
BATCH_SIZE = 500
FILE_ERRORS = (OSError, BadZipFile, InvalidFileException, KeyError, ValueError)


@dataclass
class ImportStats:
    processed: int = 0
    created: int = 0
    skipped: int = 0
    errors: int = 0


class WorkbookError(ValueError):
    pass


class RowError(ValueError):
    pass


def _parse_headers(row) -> dict[str, int]:
    indexes: dict[str, int] = {}
    for index, cell in enumerate(row):
        name = str(cell.value).strip() if cell.value is not None else ""
        if name in COLUMNS and name not in indexes:
            indexes[name] = index
    missing = set(COLUMNS) - indexes.keys()
    if missing:
        raise WorkbookError(f"В файле отсутствуют колонки: {', '.join(sorted(missing))}.")
    return indexes


def _parse_row(row, indexes: dict[str, int]) -> Mailing:
    values: dict[str, object] = {}
    for name, index in indexes.items():
        value = row[index].value if index < len(row) else None
        if value is None or (isinstance(value, str) and not value.strip()):
            raise RowError(f"Поле {name} обязательно для заполнения.")
        values[name] = value

    external_id = values["external_id"]
    if isinstance(external_id, float) and external_id.is_integer():
        external_id = int(external_id)
    external_id = str(external_id).strip()

    user_id = values["user_id"]
    try:
        user_id = int(str(user_id).strip())
    except ValueError as exc:
        raise RowError("user_id должен быть целым числом.") from exc
    if user_id <= 0:
        raise RowError("user_id должен быть положительным числом.")

    email = str(values["email"]).strip()
    try:
        validate_email(email)
    except ValidationError as exc:
        raise RowError(f"Некорректный email: {email!r}.") from exc
    if len(email) > Mailing._meta.get_field("email").max_length:
        raise RowError("email превышает допустимую длину.")

    subject = str(values["subject"]).strip()
    if len(subject) > Mailing._meta.get_field("subject").max_length:
        raise RowError("Тема письма превышает допустимую длину.")

    message = str(values["message"]).strip()

    return Mailing(
        external_id=external_id,
        user_id=user_id,
        email=email,
        subject=subject,
        message=message,
    )


def _save_batch(batch: list[Mailing], stats: ImportStats) -> None:
    if not batch:
        return
    # Keep the first occurrence when the same external_id repeats within a batch.
    unique: dict[str, Mailing] = {}
    for mailing in batch:
        unique.setdefault(mailing.external_id, mailing)
    stats.skipped += len(batch) - len(unique)

    # Per-row atomic insert-or-skip: UNIQUE alone decides, so a concurrent
    # insert of the same external_id can't be double-counted.
    for mailing in unique.values():
        try:
            with transaction.atomic():
                mailing.save(force_insert=True)
        except IntegrityError:
            stats.skipped += 1
        else:
            stats.created += 1
    batch.clear()


def import_mailings(path: Path, batch_size: int = BATCH_SIZE) -> ImportStats:
    """Stream an XLSX file's first worksheet into the Mailing queue."""
    if path.suffix.lower() != ".xlsx":
        raise WorkbookError("Ожидается файл с расширением .xlsx.")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except FILE_ERRORS as exc:
        raise WorkbookError("Не удалось прочитать XLSX-файл.") from exc

    stats = ImportStats()
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows()
        indexes = _parse_headers(next(rows, ()))

        batch: list[Mailing] = []
        for number, row in enumerate(rows, start=2):
            if all(cell.value is None for cell in row):
                continue
            stats.processed += 1
            try:
                batch.append(_parse_row(row, indexes))
            except RowError as exc:
                stats.errors += 1
                logger.warning("Строка %s: %s", number, exc)
                continue
            if len(batch) >= batch_size:
                _save_batch(batch, stats)
        _save_batch(batch, stats)
    finally:
        workbook.close()
    return stats
